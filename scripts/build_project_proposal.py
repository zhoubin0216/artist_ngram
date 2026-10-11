from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "outputs" / "Project_Proposal.docx"
FONT = "Arial"


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_style_font(style, size, bold=False, color=(0, 0, 0)):
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(*color)


def add_body_paragraph(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.widow_control = True
    run = p.add_run(text)
    set_run_font(run, 9.5)
    return p


def add_bullet(doc, label, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    lead = p.add_run(label)
    set_run_font(lead, 9.3, bold=True)
    body = p.add_run(text)
    set_run_font(body, 9.3)
    return p


def add_step(doc, number, label, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    n = p.add_run(f"{number}. ")
    set_run_font(n, 9.3, bold=True)
    lead = p.add_run(label)
    set_run_font(lead, 9.3, bold=True)
    body = p.add_run(text)
    set_run_font(body, 9.3)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.60)
section.bottom_margin = Inches(0.60)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
set_style_font(styles["Normal"], 9.5)
styles["Normal"].paragraph_format.space_after = Pt(4)

set_style_font(styles["Title"], 19, bold=True)
styles["Title"].paragraph_format.space_after = Pt(4)
styles["Title"].paragraph_format.keep_with_next = True

set_style_font(styles["Subtitle"], 9.5, color=(75, 85, 99))
styles["Subtitle"].paragraph_format.space_after = Pt(8)
styles["Subtitle"].paragraph_format.keep_with_next = True

set_style_font(styles["Heading 1"], 11.5, bold=True)
styles["Heading 1"].paragraph_format.space_before = Pt(8)
styles["Heading 1"].paragraph_format.space_after = Pt(3)
styles["Heading 1"].paragraph_format.keep_with_next = True

for style_name in ["Title", "Subtitle", "Heading 1"]:
    p_pr = styles[style_name]._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = title.add_run("European Artist Visibility in English Books")
set_run_font(run, 19, bold=True)

subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = subtitle.add_run("Data Intensive Computing Project Proposal")
set_run_font(run, 9.5, italic=True, color=(75, 85, 99))

add_body_paragraph(
    doc,
    "This project will use Google Books N-gram records and a curated artist dictionary to measure how the visibility of European painters and art movements changed in English-language books. Spark will produce normalized yearly trends and visual comparisons.",
)

doc.add_paragraph("1 Problem Description", style="Heading 1")
add_body_paragraph(
    doc,
    "The project will investigate how often major European artists were mentioned from 1800 to the latest year available in the selected corpus. It will identify the most visible artists and movements, changes in attention over time, and peaks during or after an artist's lifetime. Because common surnames and spelling variants can create false matches, the first version will use complete English names and authority-linked records to prioritize precision.",
)

doc.add_paragraph("2 Tools", style="Heading 1")
add_bullet(doc, "HDFS or equivalent storage. ", "Store compressed N-gram files, the artist dictionary, and intermediate data.")
add_bullet(doc, "Apache Spark with PySpark. ", "Filter files in parallel, join metadata, and aggregate yearly counts.")
add_bullet(doc, "Parquet. ", "Store the filtered and aggregated results efficiently.")
add_bullet(doc, "Plotly or Matplotlib. ", "Create time-series charts, movement comparisons, rankings, and heat maps.")

doc.add_paragraph("3 Data", style="Heading 1")
add_body_paragraph(
    doc,
    "The main source will be the English Google Books N-gram dataset. The project will use the files needed for multi-word names and yearly corpus totals for normalization. The analysis period will run from 1800 to the latest year supported by the selected release.",
)
add_body_paragraph(
    doc,
    "The lookup dataset contains 1,370 European painters across 15 movements, from Neoclassicism to early abstract art. Each record includes a Wikidata QID, canonical English name, birth and death years, nationality, movement, and Getty ULAN ID. The artist data was collected from Wikidata and checked against Getty ULAN. Raw N-gram files will be downloaded, stored in HDFS, filtered with Spark, and saved as smaller Parquet files.",
)

doc.add_paragraph("4 Methodology and Algorithm", style="Heading 1")
add_step(doc, "1", "Prepare the dictionary. ", "Keep unique full names, standardize spacing and capitalization, and remove one-token or duplicate names.")
add_step(doc, "2", "Filter the corpus. ", "Broadcast the small query dictionary in Spark and retain exact matches from relevant N-gram partitions.")
add_step(doc, "3", "Aggregate yearly attention. ", "Sum counts by artist and year, then join birth, death, nationality, and movement metadata.")
add_step(doc, "4", "Normalize counts. ", "Divide artist counts by the corpus total for each year to control for changes in corpus size.")
add_step(doc, "5", "Check ambiguity. ", "Flag names with substantial frequency before birth or active years for exclusion or review.")
add_step(doc, "6", "Visualize results. ", "Plot artist trends, compare movements, rank peak and cumulative visibility, and apply a short moving average.")

add_body_paragraph(
    doc,
    "The result will be a reproducible pipeline and a concise set of visualizations showing how the English-language visibility of European artists changed over time.",
)

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True
    for run in paragraph.runs:
        if run.font.name is None:
            set_run_font(run, 9.5)

doc.core_properties.title = "European Artist Visibility in English Books"
doc.core_properties.subject = "Data Intensive Computing Project Proposal"
doc.core_properties.keywords = "Google Books N-gram, European artists, Spark, HDFS, Wikidata"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
